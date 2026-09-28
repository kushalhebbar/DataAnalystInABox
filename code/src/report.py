from __future__ import annotations

from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from .state import RunState


def _run_dir(state: RunState) -> Path:
    d = Path("artifacts") / state.config.run_id
    d.mkdir(parents=True, exist_ok=True)
    return d


def build_pdf(state: RunState, output_path: str | None = None) -> str:
    path = Path(output_path) if output_path else _run_dir(state) / "report.pdf"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Caption", parent=styles["Italic"], fontSize=8, textColor=colors.grey))

    doc = SimpleDocTemplate(str(path), pagesize=LETTER, title=f"Analysis Report {state.config.run_id}")
    flow = []

    flow.append(Paragraph("Data Analysis Report", styles["Title"]))
    flow.append(Paragraph(f"Run {state.config.run_id} · {state.config.created_at}", styles["Caption"]))
    flow.append(Spacer(1, 0.2 * inch))

    problem = state.improved_problem_statement or state.inputs.problem_statement
    flow.append(Paragraph("Problem", styles["Heading2"]))
    flow.append(Paragraph(problem, styles["BodyText"]))
    flow.append(Paragraph(f"Audience: {state.inputs.audience}", styles["Caption"]))
    flow.append(Spacer(1, 0.15 * inch))

    if state.insights:
        flow.append(Paragraph("Executive Summary", styles["Heading2"]))
        flow.append(Paragraph(f"<b>{state.insights.headline}</b>", styles["BodyText"]))
        flow.append(Spacer(1, 0.05 * inch))
        flow.append(Paragraph(state.insights.narrative, styles["BodyText"]))
        if state.insights.recommendations:
            flow.append(Paragraph("Recommendations", styles["Heading3"]))
            flow.append(_bullets(state.insights.recommendations, styles))
        if state.insights.risks:
            flow.append(Paragraph("Risks & Caveats", styles["Heading3"]))
            flow.append(_bullets(state.insights.risks, styles))
        flow.append(Spacer(1, 0.15 * inch))

    if state.eda and state.eda.findings:
        flow.append(Paragraph("Key Findings", styles["Heading2"]))
        for f in state.eda.findings:
            flow.append(Paragraph(f"• {f.statement}", styles["BodyText"]))
            flow.append(Paragraph(f"{f.evidence} — confidence {int(f.confidence * 100)}%. {f.caveat}", styles["Caption"]))
            flow.append(Spacer(1, 0.05 * inch))

    if state.eda and state.eda.charts:
        flow.append(Paragraph("Charts", styles["Heading2"]))
        for chart in state.eda.charts:
            if not Path(chart.path).exists():
                continue
            flow.append(Paragraph(chart.title, styles["Heading4"]))
            flow.append(Image(chart.path, width=4.5 * inch, height=3.0 * inch, kind="proportional"))
            if chart.caption:
                flow.append(Paragraph(chart.caption, styles["Caption"]))
            flow.append(Spacer(1, 0.15 * inch))

    doc.build(flow)
    return str(path)


def _bullets(items: list[str], styles) -> ListFlowable:
    return ListFlowable(
        [ListItem(Paragraph(item, styles["BodyText"])) for item in items],
        bulletType="bullet",
    )


def build_pptx(state: RunState, output_path: str | None = None) -> str:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    path = Path(output_path) if output_path else _run_dir(state) / "report.pptx"
    prs = Presentation()
    blank = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]

    title_slide = prs.slides.add_slide(title_layout)
    title_slide.shapes.title.text = "Data Analysis Report"
    title_slide.placeholders[1].text = f"Run {state.config.run_id} · {state.inputs.audience}"

    if state.insights:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Executive Summary"
        body = slide.placeholders[1].text_frame
        body.text = state.insights.headline
        for line in state.insights.narrative.split(". "):
            line = line.strip()
            if line:
                p = body.add_paragraph()
                p.text = line
                p.font.size = Pt(14)

    if state.eda and state.eda.findings:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Key Findings"
        tf = slide.placeholders[1].text_frame
        tf.text = state.eda.findings[0].statement
        for f in state.eda.findings[1:]:
            p = tf.add_paragraph()
            p.text = f.statement
            p.font.size = Pt(14)

    for chart in (state.eda.charts if state.eda else []):
        if not Path(chart.path).exists():
            continue
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(chart.path, Inches(0.5), Inches(0.6), width=Inches(9))

    prs.save(str(path))
    return str(path)
